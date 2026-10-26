# -*- coding: utf-8 -*-
"""把"本机硬件/软件环境"表插入报告 4.1 配置小节（放在训练配置表之后）。

为什么单独写脚本而不是重跑 pandoc：
    当前 report_polished.docx 已经把字号、行距、三线表、页边标记都修好了，
    重跑 pandoc 会把这些全部推倒重来。这里只在现有文档上做定点插入。
    如果文档里已经有这段（重复运行），直接跳过，保证可重复执行。

用法：
    python add_env_block.py --doc ..\report_polished.docx
"""
import argparse

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from polish_docx import apply_run_font, three_line

LEAD = "实验环境（全部训练与评测均在本机单卡完成）"
ROWS = [
    ("项目", "配置"),
    ("CPU", "AMD Ryzen 7 7745HX，8 核 16 线程"),
    ("内存", "32 GB DDR5-5600（2×16 GB）"),
    ("GPU", "NVIDIA GeForce RTX 4060 Laptop，8 GB 显存，驱动 610.47"),
    ("操作系统", "Windows 11 家庭版 中文版（10.0.26300）"),
    ("软件栈", "Python 3.11.4 / PyTorch 2.14.1+cu126 / torchvision 0.29.1 / CUDA 12.6 / cuDNN 9.10.2"),
    ("训练峰值显存", "ResNet-50 batch 4 @513 约 3.3 GB；ResNet-101 约 5.2 GB（故训练时不并行评测）"),
    ("存储", "数据集、训练输出与提交物在 E 盘；验证集概率缓存 17.8 GB 亦置于 E 盘"),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--doc", required=True)
    args = ap.parse_args()

    d = Document(args.doc)
    if any(LEAD in p.text for p in d.paragraphs):
        print("已存在实验环境段，跳过")
        return

    anchor = None
    for t in d.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "配置项":
            anchor = t
            break
    if anchor is None:
        raise SystemExit("找不到 4.1 的训练配置表（首列表头应为“配置项”）")

    # 引导句：与正文同字号，加粗
    p = d.add_paragraph()
    run = p.add_run(LEAD)
    apply_run_font(run, "Times New Roman", "宋体", 12, bold=True)
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.5

    table = d.add_table(rows=len(ROWS), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(ROWS):
        for j, text in enumerate((k, v)):
            cell = table.rows[i].cells[j]
            cell.width = Cm(3.2 if j == 0 else 12.8)
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.line_spacing = 1.15
            cp.paragraph_format.space_before = Pt(1)
            cp.paragraph_format.space_after = Pt(1)
            r = cp.add_run(text)
            apply_run_font(r, "Times New Roman", "宋体", 11, bold=(i == 0))
    three_line(table)

    # 定点移动到训练配置表之后：anchor_tbl -> 引导句 -> 新表
    anchor._tbl.addnext(p._p)
    p._p.addnext(table._tbl)

    d.save(args.doc)
    print(f"已插入实验环境表（{len(ROWS) - 1} 项）-> {args.doc}")


if __name__ == "__main__":
    main()
